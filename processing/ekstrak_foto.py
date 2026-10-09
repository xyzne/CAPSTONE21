"""Ambil foto dokumentasi kebun dari Excel (gambar di kolom Dokumentasi) -> data/foto/kebun-NN.jpg
Jalankan: python ekstrak_foto.py <Jenis_Kebun.xlsx> data
Nomor pada nama file = kolom No pada baris tempat gambar berada. Butuh: openpyxl"""
import sys, os
from openpyxl import load_workbook
XLSX, OUT = sys.argv[1], sys.argv[2]
ws = load_workbook(XLSX).active
os.makedirs(f"{OUT}/foto", exist_ok=True)
n = 0
for im in ws._images:
    no = ws.cell(row=im.anchor._from.row + 1, column=1).value
    open(f"{OUT}/foto/kebun-{int(no):02d}.jpg", "wb").write(im._data()); n += 1
print(n, "foto disimpan")
